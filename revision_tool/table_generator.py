# Copyright (c) 2026 王刚. All rights reserved.
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn as docx_qn

from .tracked_changes import Revision


_TABLE_TITLE = "制度修订对照表"
_HEADER_ROW = ["序号", "修订条款", "修订前内容", "修订后内容"]
_MAX_CELL_CHARS = 500


def _set_cell_text(cell, text: str, bold: bool = False, font_size: float = 9,
                   alignment=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.15

    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(docx_qn("w:eastAsia"), "宋体")


def _set_cell_shading(cell, color: str):
    shading = cell._element.get_or_add_tcPr()
    shd_elem = shading.makeelement(docx_qn("w:shd"), {
        docx_qn("w:val"): "clear",
        docx_qn("w:color"): "auto",
        docx_qn("w:fill"): color,
    })
    shading.append(shd_elem)


def _set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else tbl._add_tblPr()
    borders = tbl_pr.makeelement(docx_qn("w:tblBorders"), {})
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.makeelement(docx_qn(f"w:{border_name}"), {
            docx_qn("w:val"): "single",
            docx_qn("w:sz"): "4",
            docx_qn("w:space"): "0",
            docx_qn("w:color"): "000000",
        })
        borders.append(border)
    existing = tbl_pr.find(docx_qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    tbl_pr.append(borders)


def _set_column_widths(table, widths: list[float]):
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def _truncate_text(text: str, max_chars: int = _MAX_CELL_CHARS) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars - 3] + "..."


def generate_comparison_table(revisions: list[Revision], doc: Document,
                              insert_position: str = "end",
                              marker_text: str = "") -> Document:
    if not revisions:
        return doc

    filtered = [r for r in revisions if r.context_before or r.context_after
                or r.deleted_text or r.inserted_text]
    if not filtered:
        return doc

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Pt(12)
    title_para.paragraph_format.space_after = Pt(6)
    run = title_para.add_run(_TABLE_TITLE)
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(docx_qn("w:eastAsia"), "黑体")

    num_rows = len(filtered) + 1
    table = doc.add_table(rows=num_rows, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)

    _set_column_widths(table, [1.5, 3.5, 6.0, 6.0])

    for col_idx, header in enumerate(_HEADER_ROW):
        cell = table.rows[0].cells[col_idx]
        _set_cell_text(cell, header, bold=True, font_size=10,
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(cell, "D9E2F3")

    for row_idx, rev in enumerate(filtered, start=1):
        seq = str(row_idx)
        clause = rev.clause_name or "—"
        before = rev.context_before if rev.context_before else (rev.deleted_text if rev.deleted_text else "（新增条款）")
        after = rev.context_after if rev.context_after else (rev.inserted_text if rev.inserted_text else "（删除条款）")
        before = _truncate_text(before)
        after = _truncate_text(after)

        _set_cell_text(table.rows[row_idx].cells[0], seq,
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(table.rows[row_idx].cells[1], clause,
                       alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(table.rows[row_idx].cells[2], before)
        _set_cell_text(table.rows[row_idx].cells[3], after)

    if insert_position != "end":
        _move_table_before_marker(doc, table, title_para, marker_text, insert_position)

    copyright_para = doc.add_paragraph()
    copyright_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    copyright_para.paragraph_format.space_before = Pt(6)
    run = copyright_para.add_run("Copyright \u00a9 2026 王刚. All rights reserved.")
    run.font.size = Pt(7)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(docx_qn("w:eastAsia"), "宋体")
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    return doc


def _move_table_before_marker(doc: Document, table, title_para,
                              marker_text: str, position: str):
    body = doc.element.body
    tbl_element = table._tbl
    title_element = title_para._element

    if position == "before_body" or position == "start":
        first_child = body[0]
        body.remove(tbl_element)
        body.remove(title_element)
        body.insert(list(body).index(first_child), title_element)
        body.insert(list(body).index(title_element) + 1, tbl_element)
    elif marker_text:
        for child in body:
            if child.tag.endswith("}p"):
                for run in child.iter():
                    if run.tag.endswith("}t") and run.text and marker_text in run.text:
                        body.remove(tbl_element)
                        body.remove(title_element)
                        idx = list(body).index(child)
                        body.insert(idx, title_element)
                        body.insert(idx + 1, tbl_element)
                        return


def generate_table_to_new_doc(revisions: list[Revision], output_path: str) -> Document:
    doc = Document()
    generate_comparison_table(revisions, doc)
    doc.save(output_path)
    return doc
