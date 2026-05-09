# Copyright (c) 2026 王刚. All rights reserved.
from pathlib import Path
from docx import Document

from .tracked_changes import extract_revisions, accept_all_revisions, Revision
from .doc_compare import compare_documents_inline
from .table_generator import generate_comparison_table, generate_table_to_new_doc


class RevisionTool:
    def __init__(self):
        self.revisions: list[Revision] = []

    def process_tracked_changes(self, doc_path: str, output_path: str | None = None,
                                accept_changes: bool = True,
                                table_position: str = "end",
                                marker_text: str = "") -> list[Revision]:
        doc = Document(doc_path)
        self.revisions = extract_revisions(doc)

        if accept_changes:
            accept_all_revisions(doc)

        if self.revisions:
            generate_comparison_table(self.revisions, doc,
                                     insert_position=table_position,
                                     marker_text=marker_text)

        save_path = output_path or _default_output_path(doc_path)
        doc.save(save_path)
        print(f"已保存至: {save_path}")
        print(f"共识别 {len(self.revisions)} 处修订")
        return self.revisions

    def process_two_docs(self, old_doc_path: str, new_doc_path: str,
                         output_path: str | None = None,
                         table_position: str = "end",
                         marker_text: str = "") -> list[Revision]:
        old_doc = Document(old_doc_path)
        new_doc = Document(new_doc_path)

        self.revisions = compare_documents_inline(old_doc, new_doc)

        if self.revisions:
            generate_comparison_table(self.revisions, new_doc,
                                     insert_position=table_position,
                                     marker_text=marker_text)

        save_path = output_path or _default_output_path(new_doc_path)
        new_doc.save(save_path)
        print(f"已保存至: {save_path}")
        print(f"共识别 {len(self.revisions)} 处修订")
        return self.revisions

    def generate_table_only(self, revisions: list[Revision] | None = None,
                            output_path: str = "revision_table.docx") -> str:
        revs = revisions or self.revisions
        if not revs:
            print("无修订记录可生成")
            return ""
        generate_table_to_new_doc(revs, output_path)
        print(f"对照表已保存至: {output_path}")
        return output_path

    def print_summary(self, revisions: list[Revision] | None = None):
        revs = revisions or self.revisions
        if not revs:
            print("无修订记录")
            return

        print(f"\n{'='*60}")
        print(f"修订记录摘要 (共 {len(revs)} 处)")
        print(f"{'='*60}")
        for i, rev in enumerate(revs, 1):
            print(f"\n[{i}] 条款: {rev.clause_name or '未识别'}")
            if rev.deleted_text:
                print(f"  删除: {rev.deleted_text[:80]}{'...' if len(rev.deleted_text) > 80 else ''}")
            if rev.inserted_text:
                print(f"  新增: {rev.inserted_text[:80]}{'...' if len(rev.inserted_text) > 80 else ''}")


def _default_output_path(original_path: str) -> str:
    p = Path(original_path)
    return str(p.parent / f"{p.stem}_含修订对照表{p.suffix}")
