import os
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from revision_tool.core import RevisionTool
from revision_tool.tracked_changes import Revision

COPYRIGHT = "Copyright \u00a9 2026 王刚. All rights reserved."
APP_VERSION = "1.0.0"


class RevisionToolApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("制度修订对照表生成工具")
        self.root.geometry("780x700")
        self.root.minsize(680, 560)
        self.tool = RevisionTool()
        self.revisions: list[Revision] = []
        self.accept_var = tk.BooleanVar(value=True)
        self.mode_var = tk.StringVar(value="tracked")
        self.position_var = tk.StringVar(value="end")
        self.marker_var = tk.StringVar()
        self.tracked_path_var = tk.StringVar()
        self.old_path_var = tk.StringVar()
        self.new_path_var = tk.StringVar()
        self.output_path_var = tk.StringVar()
        self.status_var = tk.StringVar(value="就绪")

        self._build_ui()
        self._center_window()

    def _center_window(self):
        self.root.update_idletasks()
        w = self.root.winfo_width()
        h = self.root.winfo_height()
        x = (self.root.winfo_screenwidth() - w) // 2
        y = (self.root.winfo_screenheight() - h) // 2
        self.root.geometry(f"+{x}+{y}")

    def _build_ui(self):
        style = ttk.Style()
        style.configure("Title.TLabel", font=("微软雅黑", 16, "bold"))
        style.configure("Section.TLabelframe.Label", font=("微软雅黑", 10, "bold"))
        style.configure("Action.TButton", font=("微软雅黑", 10))

        main_frame = ttk.Frame(self.root, padding=15)
        main_frame.pack(fill=tk.BOTH, expand=True)

        title_label = ttk.Label(main_frame, text="制度修订对照表生成工具", style="Title.TLabel")
        title_label.pack(pady=(0, 15))

        mode_frame = ttk.LabelFrame(main_frame, text="处理模式", style="Section.TLabelframe", padding=10)
        mode_frame.pack(fill=tk.X, pady=(0, 10))

        ttk.Radiobutton(mode_frame, text="模式一：处理带修订痕迹的单个文档",
                        variable=self.mode_var, value="tracked",
                        command=self._on_mode_change).pack(anchor=tk.W, pady=2)
        ttk.Radiobutton(mode_frame, text="模式二：对比原文档与修订后文档",
                        variable=self.mode_var, value="compare",
                        command=self._on_mode_change).pack(anchor=tk.W, pady=2)

        file_frame = ttk.LabelFrame(main_frame, text="文件选择", style="Section.TLabelframe", padding=10)
        file_frame.pack(fill=tk.X, pady=(0, 10))

        self.tracked_frame = ttk.Frame(file_frame)
        self.tracked_frame.pack(fill=tk.X)

        ttk.Label(self.tracked_frame, text="修订文档：").grid(row=0, column=0, sticky=tk.W, pady=4)
        ttk.Entry(self.tracked_frame, textvariable=self.tracked_path_var, width=55).grid(row=0, column=1, padx=5, pady=4)
        ttk.Button(self.tracked_frame, text="浏览...", command=lambda: self._browse_file(self.tracked_path_var, "选择修订文档")).grid(row=0, column=2, pady=4)

        self.compare_frame = ttk.Frame(file_frame)

        ttk.Label(self.compare_frame, text="原文档：").grid(row=0, column=0, sticky=tk.W, pady=4)
        ttk.Entry(self.compare_frame, textvariable=self.old_path_var, width=55).grid(row=0, column=1, padx=5, pady=4)
        ttk.Button(self.compare_frame, text="浏览...", command=lambda: self._browse_file(self.old_path_var, "选择原文档")).grid(row=0, column=2, pady=4)

        ttk.Label(self.compare_frame, text="修订后文档：").grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Entry(self.compare_frame, textvariable=self.new_path_var, width=55).grid(row=1, column=1, padx=5, pady=4)
        ttk.Button(self.compare_frame, text="浏览...", command=lambda: self._browse_file(self.new_path_var, "选择修订后文档")).grid(row=1, column=2, pady=4)

        self._on_mode_change()

        options_frame = ttk.LabelFrame(main_frame, text="选项", style="Section.TLabelframe", padding=10)
        options_frame.pack(fill=tk.X, pady=(0, 10))

        opt_row1 = ttk.Frame(options_frame)
        opt_row1.pack(fill=tk.X, pady=2)

        ttk.Checkbutton(opt_row1, text="接受所有修订", variable=self.accept_var).pack(side=tk.LEFT, padx=(0, 20))

        ttk.Label(opt_row1, text="对照表位置：").pack(side=tk.LEFT)
        pos_combo = ttk.Combobox(opt_row1, textvariable=self.position_var,
                                 values=["end", "start", "marker"], state="readonly", width=10)
        pos_combo.pack(side=tk.LEFT, padx=5)
        pos_combo.bind("<<ComboboxSelected>>", self._on_position_change)

        self.marker_label = ttk.Label(opt_row1, text="标记文本：")
        self.marker_label.pack(side=tk.LEFT, padx=(15, 0))
        self.marker_entry = ttk.Entry(opt_row1, textvariable=self.marker_var, width=20, state=tk.DISABLED)
        self.marker_entry.pack(side=tk.LEFT, padx=5)

        output_frame = ttk.Frame(options_frame)
        output_frame.pack(fill=tk.X, pady=(8, 0))

        ttk.Label(output_frame, text="输出路径：").pack(side=tk.LEFT)
        ttk.Entry(output_frame, textvariable=self.output_path_var, width=50).pack(side=tk.LEFT, padx=5)
        ttk.Button(output_frame, text="浏览...", command=self._browse_output).pack(side=tk.LEFT)

        action_frame = ttk.Frame(main_frame)
        action_frame.pack(fill=tk.X, pady=(5, 10))

        self.run_btn = ttk.Button(action_frame, text="生成修订对照表", style="Action.TButton",
                                  command=self._run)
        self.run_btn.pack(side=tk.LEFT, padx=(0, 10))

        self.table_only_btn = ttk.Button(action_frame, text="仅导出对照表", command=self._export_table_only)
        self.table_only_btn.pack(side=tk.LEFT, padx=(0, 10))

        self.progress = ttk.Progressbar(action_frame, mode="indeterminate", length=200)
        self.progress.pack(side=tk.LEFT, padx=5)

        result_frame = ttk.LabelFrame(main_frame, text="修订记录", style="Section.TLabelframe", padding=10)
        result_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 5))

        columns = ("seq", "clause", "before", "after")
        self.tree = ttk.Treeview(result_frame, columns=columns, show="headings", height=8)
        self.tree.heading("seq", text="序号")
        self.tree.heading("clause", text="修订条款")
        self.tree.heading("before", text="修订前内容")
        self.tree.heading("after", text="修订后内容")
        self.tree.column("seq", width=45, anchor=tk.CENTER, stretch=False)
        self.tree.column("clause", width=120, anchor=tk.CENTER, stretch=False)
        self.tree.column("before", width=250, stretch=True)
        self.tree.column("after", width=250, stretch=True)

        scrollbar_y = ttk.Scrollbar(result_frame, orient=tk.VERTICAL, command=self.tree.yview)
        scrollbar_x = ttk.Scrollbar(result_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        scrollbar_y.grid(row=0, column=1, sticky="ns")
        scrollbar_x.grid(row=1, column=0, sticky="ew")
        result_frame.columnconfigure(0, weight=1)
        result_frame.rowconfigure(0, weight=1)

        status_bar = ttk.Label(main_frame, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(fill=tk.X, pady=(5, 0))

        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill=tk.X, pady=(2, 0))

        copyright_label = ttk.Label(bottom_frame, text=COPYRIGHT, foreground="gray",
                                    font=("微软雅黑", 8))
        copyright_label.pack(side=tk.LEFT)

        about_btn = ttk.Button(bottom_frame, text="关于", command=self._show_about)
        about_btn.pack(side=tk.RIGHT)

    def _on_mode_change(self, event=None):
        mode = self.mode_var.get()
        if mode == "tracked":
            self.compare_frame.pack_forget()
            self.tracked_frame.pack(fill=tk.X)
            self.accept_var.set(True)
        else:
            self.tracked_frame.pack_forget()
            self.compare_frame.pack(fill=tk.X)

    def _on_position_change(self, event=None):
        if self.position_var.get() == "marker":
            self.marker_entry.config(state=tk.NORMAL)
        else:
            self.marker_entry.config(state=tk.DISABLED)
            self.marker_var.set("")

    def _browse_file(self, path_var: tk.StringVar, title: str):
        filetypes = [("Word 文档", "*.docx"), ("所有文件", "*.*")]
        path = filedialog.askopenfilename(title=title, filetypes=filetypes)
        if path:
            path_var.set(path)
            if not self.output_path_var.get():
                self._auto_set_output(path)

    def _browse_output(self):
        filetypes = [("Word 文档", "*.docx"), ("所有文件", "*.*")]
        path = filedialog.asksaveasfilename(title="保存为", filetypes=filetypes,
                                            defaultextension=".docx")
        if path:
            self.output_path_var.set(path)

    def _auto_set_output(self, input_path: str):
        p = Path(input_path)
        output = str(p.parent / f"{p.stem}_含修订对照表{p.suffix}")
        self.output_path_var.set(output)

    def _validate_inputs(self) -> str | None:
        mode = self.mode_var.get()
        if mode == "tracked":
            path = self.tracked_path_var.get().strip()
            if not path:
                return "请选择修订文档"
            if not os.path.isfile(path):
                return f"文件不存在: {path}"
        else:
            old = self.old_path_var.get().strip()
            new = self.new_path_var.get().strip()
            if not old:
                return "请选择原文档"
            if not new:
                return "请选择修订后文档"
            if not os.path.isfile(old):
                return f"文件不存在: {old}"
            if not os.path.isfile(new):
                return f"文件不存在: {new}"
        return None

    def _run(self):
        error = self._validate_inputs()
        if error:
            messagebox.showwarning("输入错误", error)
            return

        self.run_btn.config(state=tk.DISABLED)
        self.table_only_btn.config(state=tk.DISABLED)
        self.progress.start(10)
        self.status_var.set("正在处理...")

        threading.Thread(target=self._do_run, daemon=True).start()

    def _do_run(self):
        try:
            mode = self.mode_var.get()
            output = self.output_path_var.get().strip() or None
            position = self.position_var.get()
            marker = self.marker_var.get().strip() if position == "marker" else ""

            if mode == "tracked":
                doc_path = self.tracked_path_var.get().strip()
                self.revisions = self.tool.process_tracked_changes(
                    doc_path=doc_path,
                    output_path=output,
                    accept_changes=self.accept_var.get(),
                    table_position=position,
                    marker_text=marker,
                )
            else:
                old_path = self.old_path_var.get().strip()
                new_path = self.new_path_var.get().strip()
                self.revisions = self.tool.process_two_docs(
                    old_doc_path=old_path,
                    new_doc_path=new_path,
                    output_path=output,
                    table_position=position,
                    marker_text=marker,
                )

            self.root.after(0, self._on_run_complete)
        except Exception as e:
            self.root.after(0, lambda: self._on_run_error(str(e)))

    def _on_run_complete(self):
        self.progress.stop()
        self.run_btn.config(state=tk.NORMAL)
        self.table_only_btn.config(state=tk.NORMAL)
        self._populate_tree()
        self.status_var.set(f"完成 — 共识别 {len(self.revisions)} 处修订")
        messagebox.showinfo("完成", f"修订对照表已生成\n共识别 {len(self.revisions)} 处修订")

    def _on_run_error(self, error_msg: str):
        self.progress.stop()
        self.run_btn.config(state=tk.NORMAL)
        self.table_only_btn.config(state=tk.NORMAL)
        self.status_var.set("处理失败")
        messagebox.showerror("错误", f"处理失败：\n{error_msg}")

    def _export_table_only(self):
        if not self.revisions:
            error = self._validate_inputs()
            if error:
                messagebox.showwarning("提示", "请先处理文档或检查输入")
                return

            mode = self.mode_var.get()
            if mode == "tracked":
                from docx import Document
                from revision_tool.tracked_changes import extract_revisions
                doc = Document(self.tracked_path_var.get().strip())
                self.revisions = extract_revisions(doc)
            else:
                from docx import Document
                from revision_tool.doc_compare import compare_documents_inline
                old_doc = Document(self.old_path_var.get().strip())
                new_doc = Document(self.new_path_var.get().strip())
                self.revisions = compare_documents_inline(old_doc, new_doc)

        if not self.revisions:
            messagebox.showinfo("提示", "未检测到修订内容")
            return

        filetypes = [("Word 文档", "*.docx"), ("所有文件", "*.*")]
        path = filedialog.asksaveasfilename(title="保存对照表", filetypes=filetypes,
                                            defaultextension=".docx")
        if path:
            self.tool.generate_table_only(self.revisions, path)
            self._populate_tree()
            self.status_var.set(f"对照表已导出 — 共 {len(self.revisions)} 处修订")

    def _populate_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        for i, rev in enumerate(self.revisions, 1):
            before = rev.context_before if rev.context_before else (rev.deleted_text if rev.deleted_text else "（新增条款）")
            after = rev.context_after if rev.context_after else (rev.inserted_text if rev.inserted_text else "（删除条款）")
            if len(before) > 100:
                before = before[:97] + "..."
            if len(after) > 100:
                after = after[:97] + "..."
            self.tree.insert("", tk.END, values=(
                i,
                rev.clause_name or "—",
                before,
                after,
            ))

    def _show_about(self):
        about_win = tk.Toplevel(self.root)
        about_win.title("关于")
        about_win.geometry("380x220")
        about_win.resizable(False, False)
        about_win.transient(self.root)
        about_win.grab_set()

        x = self.root.winfo_x() + (self.root.winfo_width() - 380) // 2
        y = self.root.winfo_y() + (self.root.winfo_height() - 220) // 2
        about_win.geometry(f"+{x}+{y}")

        frame = ttk.Frame(about_win, padding=20)
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frame, text="制度修订对照表生成工具",
                  font=("微软雅黑", 14, "bold")).pack(pady=(0, 5))
        ttk.Label(frame, text=f"版本 {APP_VERSION}",
                  font=("微软雅黑", 10)).pack(pady=(0, 10))
        ttk.Separator(frame, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=5)
        ttk.Label(frame, text=COPYRIGHT,
                  font=("微软雅黑", 9), foreground="gray").pack(pady=(10, 5))
        ttk.Label(frame, text="未经授权不得复制、修改或分发本软件。",
                  font=("微软雅黑", 8), foreground="gray").pack(pady=(0, 10))
        ttk.Button(frame, text="确定", command=about_win.destroy, width=10).pack()


def run_gui():
    root = tk.Tk()
    app = RevisionToolApp(root)
    root.mainloop()


if __name__ == "__main__":
    run_gui()
